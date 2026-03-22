import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from garmin_exporter import activity_is_in_range, activity_to_row, parse_activity_date, write_excel


class GarminExporterTests(unittest.TestCase):
    def test_parse_activity_date_with_space_separator(self) -> None:
        activity = {"startTimeLocal": "2025-02-10 07:15:30"}
        self.assertEqual(parse_activity_date(activity), date(2025, 2, 10))

    def test_activity_range_filter(self) -> None:
        activity = {"startTimeLocal": "2025-02-10 07:15:30"}
        self.assertTrue(
            activity_is_in_range(activity, from_date=date(2025, 2, 1), to_date=date(2025, 2, 28))
        )
        self.assertFalse(
            activity_is_in_range(activity, from_date=date(2025, 3, 1), to_date=date(2025, 3, 30))
        )

    def test_activity_to_row_maps_expected_fields(self) -> None:
        activity = {
            "activityId": 123,
            "activityName": "Morning Run",
            "activityType": {"typeKey": "running"},
            "distance": 5000.0,
        }
        row = activity_to_row(activity)
        self.assertEqual(row["activity_id"], 123)
        self.assertEqual(row["activity_name"], "Morning Run")
        self.assertEqual(row["activity_type"], "running")
        self.assertEqual(row["distance_m"], 5000.0)

    def test_write_excel_writes_one_row_per_activity(self) -> None:
        rows = [
            {"activity_id": 1, "activity_name": "A"},
            {"activity_id": 2, "activity_name": "B"},
            {"activity_id": 3, "activity_name": "C"},
        ]
        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "activities.xlsx"
            write_excel(rows, output_path)

            workbook = load_workbook(output_path)
            worksheet = workbook["Activities"]
            # Header + 3 activity rows.
            self.assertEqual(worksheet.max_row, 4)
            self.assertEqual(worksheet["A2"].value, 1)
            self.assertEqual(worksheet["A3"].value, 2)
            self.assertEqual(worksheet["A4"].value, 3)


if __name__ == "__main__":
    unittest.main()
